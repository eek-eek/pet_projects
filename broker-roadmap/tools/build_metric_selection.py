#!/usr/bin/env python3
"""Форма выбора метрик из стандартного набора отрасли (этап 1 дорожной карты).

Собирает 01b_best_practice_metrics.xlsx из таблиц 01b_best_practice_metrics.md:
лист «Метрики» — список метрик с колонкой «Выбор», лист «Вопросы» — вопросы
для закрытия этапа 1. Текст метрик правится в markdown, затем форма
пересобирается.

Если вы уже заполнили форму, генератор не перезапишет файл без флага --force.

Запуск:  python3 tools/build_metric_selection.py [--out PATH] [--force]
После сборки откройте файл в Excel или пересчитайте через LibreOffice,
чтобы счётчик отметок получил значение.
"""

import argparse
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import build_model as bm
from build_model import BLUE, BOX, FILL_INPUT, FILL_RESULT, GREY, THIN, WRAP, _fill, font, header_row, widths

ROOT = Path(__file__).resolve().parent.parent          # broker-roadmap/
SOURCE = ROOT / "01b_best_practice_metrics.md"
DEFAULT_OUT = ROOT / "01b_best_practice_metrics.xlsx"

# Коды, на которые могут ссылаться метрики: показатели модели и параметры листа «Вводные».
MODEL_CODES = (set(bm.INPUTS) | set(bm.CALC) | {e[0] for e in bm.EXTRA} | {g[0] for g in bm.GUARDRAILS}
               | {x[0] for x in bm.EXTERNAL} | {"SHOCK", "HORIZON", "RATE_SHIFT", "PT_MRG", "PT_FLOAT"})

CHOICES = ["Да", "Нет", "Обсудить"]
CHOICE_FILLS = {"Да": FILL_RESULT, "Обсудить": _fill("FFFFF2CC"), "Нет": _fill("FFF2F2F2")}
HEAD_ROW = 6
COLUMNS = [  # заголовок, ширина
    ("№", 6), ("Группа", 16), ("Метрика", 32), ("Как считается", 32), ("Зачем", 30),
    ("Кто раскрывает / рамка", 26), ("Рекомендую в базовый набор", 14), ("Есть в модели этапа 1", 24),
    ("Выбор", 11), ("Комментарий", 28),
]
CHOICE_COL, NOTE_COL = "I", "J"
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
GROUP_TOP = Border(left=THIN, right=THIN, bottom=THIN, top=Side(style="medium", color="FF1F3864"))

QUESTIONS = [  # номер, вопрос, варианты (без запятых — это список проверки данных), рекомендация, почему
    ("В1", "Какая метрика станет North Star — главной метрикой ценности для клиента и бизнеса?",
     ["2.2 Чистый приток", "1.1 Фондированные клиенты", "2.1 Активы клиентов", "Другая — в комментарии"],
     "2.2 Чистый приток",
     "Приток измеряется в деньгах, поэтому сопоставим между сегментами — в отличие от числа клиентов. "
     "Он не зависит от движения рынка — в отличие от активов. За притоком растут активы, а с ними процентный "
     "доход и комиссии. Ограничение: у институционалов выручку больше определяет оборот, чем активы."),
    ("В2", "Что ставим наверх дерева метрик?",
     ["Оба уровня: маржа до налогов и cost/income наверху; маржинальный доход для гипотез",
      "Маржинальный доход — как сейчас", "Маржа до налогов и cost/income вместо маржинального дохода"],
     "Оба уровня",
     "Для управления и сравнения с рынком — маржа до налогов и cost/income, как принято в отрасли. Гипотезы "
     "оцениваем приростом маржинального дохода, а затраты проектов учтём на этапе 3. Для верхнего уровня "
     "в модель нужно добавить постоянные расходы: ИТ, персонал, офис."),
    ("В3", "Что делаем с расчётной моделью этапа 1?",
     ["Калькулятор под выбранным набором", "Только выбранный набор без модели"],
     "Калькулятор под выбранным набором",
     "Модель нужна, чтобы на этапе 2 оценивать эффект гипотез в деньгах. Её показатели переименую "
     "в отраслевые термины и добавлю выбранные метрики, которых в ней нет."),
]


# --- Чтение markdown ------------------------------------------------------------
def parse_metrics(text):
    """Строки таблиц групп: [номер, группа, метрика, как считается, зачем, кто раскрывает, в модели]."""
    rows, group = [], None
    for line in text.splitlines():
        m = re.match(r"^## (\d+)\. (.+)$", line)
        if m:
            group = f"{m.group(1)}. {m.group(2)}"
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if not (line.startswith("| ") and re.fullmatch(r"\d+\.\d+", cells[0])):
            continue
        assert len(cells) == 6, f"{cells[0]}: в таблице метрик должно быть 6 колонок"
        unknown = set(re.findall(r"`([A-Z][A-Z0-9_]*)`", cells[5])) - MODEL_CODES
        assert not unknown, f"{cells[0]}: в модели нет кодов {sorted(unknown)}"
        rows.append([cells[0], group] + [c.replace("`", "") for c in cells[1:]])
    return rows


def parse_base_set(text):
    """Номер метрики → место в базовом наборе из раздела 10: «1» или «4 — или 3.1»."""
    section = text.split("\n## 10.", 1)[1].split("\n## ", 1)[0]
    base = {}
    for line in section.splitlines():
        m = re.match(r"^(\d+)\. .+ — (.+)\.$", line)
        if not m:
            continue
        numbers = re.findall(r"\d+\.\d+", m.group(2))
        for n in numbers:
            other = [o for o in numbers if o != n]
            base[n] = f"{m.group(1)} — или {other[0]}" if " или " in m.group(2) else m.group(1)
    return base


# --- Листы ------------------------------------------------------------------------
def note_row(ws, row, text, height):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(COLUMNS))
    c = ws.cell(row=row, column=1, value=text)
    c.font = font()
    c.alignment = WRAP
    ws.row_dimensions[row].height = height


def build_metrics(ws, rows, base):
    first, last = HEAD_ROW + 1, HEAD_ROW + len(rows)
    choices = f"${CHOICE_COL}${first}:${CHOICE_COL}${last}"
    ws.title = "Метрики"
    ws["A1"] = "Метрики брокерского бизнеса: выбор из лучших практик"
    ws["A1"].font = font(bold=True, size=14)
    note_row(ws, 2, "Как заполнить: в жёлтой колонке «Выбор» отметьте «Да», «Нет» или «Обсудить», комментарий — "
                    "по желанию. Например: 2.2 «Чистый приток» → «Да», комментарий «нужен в разрезе сегментов». "
                    "Верните файл или пришлите номера метрик. Ещё три вопроса — на листе «Вопросы».", 28)
    note_row(ws, 3, "«Рекомендую в базовый набор» — номер в рекомендуемом наборе из 10 ключевых метрик: одинаковый "
                    "номер — одна метрика набора, «или» — достаточно одной из двух. «Есть в модели»: да — считается; "
                    "выводится — простой формулой из показателей модели; паспорт — описана как опережающий индикатор, "
                    "в расчёте не участвует; ограничитель — описана как ограничитель; частично — в другом разрезе; "
                    "нет — добавим, если выберете.", 40)
    ws["A4"] = (f'="Отмечено: да — "&COUNTIF({choices},"Да")&" · обсудить — "&COUNTIF({choices},"Обсудить")'
                f'&" · нет — "&COUNTIF({choices},"Нет")&" · без отметки — "&COUNTBLANK({choices})'
                f'&" из "&ROWS({choices})')
    ws["A4"].font = font(bold=True)

    header_row(ws, HEAD_ROW, [h for h, _ in COLUMNS], height=42)
    widths(ws, {get_column_letter(i): w for i, (_, w) in enumerate(COLUMNS, start=1)})
    prev_group = rows[0][1]
    for r, (number, group, name, how, why, who, model) in enumerate(rows, start=first):
        values = [number, group, name, how, why, who, base.get(number, ""), model, None, None]
        for col, value in enumerate(values, start=1):
            c = ws.cell(row=r, column=col, value=value)
            c.font = font(GREY) if col == 8 else font(BLUE) if col >= 9 else font(bold=col == 3)
            c.alignment = CENTER if col in (1, 7, 9) else WRAP
            c.border = GROUP_TOP if group != prev_group else BOX
            if col >= 9:
                c.fill = FILL_INPUT
        prev_group = group

    dv = DataValidation(type="list", formula1='"' + ",".join(CHOICES) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Выбор", error="Выберите «Да», «Нет» или «Обсудить».")
    ws.add_data_validation(dv)
    dv.add(f"{CHOICE_COL}{first}:{CHOICE_COL}{last}")
    for choice, fill in CHOICE_FILLS.items():
        ws.conditional_formatting.add(
            f"A{first}:{NOTE_COL}{last}",
            FormulaRule(formula=[f'${CHOICE_COL}{first}="{choice}"'], fill=fill,
                        font=Font(color=GREY) if choice == "Нет" else None))
    ws.auto_filter.ref = f"A{HEAD_ROW}:{NOTE_COL}{last}"
    ws.freeze_panes = f"D{first}"
    ws.print_title_rows = f"{HEAD_ROW}:{HEAD_ROW}"


def build_questions(ws):
    ws["A1"] = "Вопросы для закрытия этапа 1"
    ws["A1"].font = font(bold=True, size=14)
    ws["A2"] = ("Выберите ответ в жёлтой колонке «Ответ» из списка или впишите свой. "
                "«Рекомендую» — моё предложение, решение за вами.")
    ws["A2"].font = font()
    header_row(ws, 4, ["№", "Вопрос", "Варианты ответа", "Рекомендую", "Почему", "Ответ", "Комментарий"])
    widths(ws, {"A": 5, "B": 32, "C": 36, "D": 20, "E": 56, "F": 30, "G": 30})
    for r, (number, question, options, recommended, why) in enumerate(QUESTIONS, start=5):
        values = [number, question, "\n".join(f"• {o}" for o in options), recommended, why, None, None]
        for col, value in enumerate(values, start=1):
            c = ws.cell(row=r, column=col, value=value)
            c.font = font(BLUE) if col >= 6 else font(bold=col == 2)
            c.alignment = WRAP
            c.border = BOX
            if col >= 6:
                c.fill = FILL_INPUT
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True,
                            showErrorMessage=False)
        ws.add_data_validation(dv)
        dv.add(f"F{r}")
    ws.freeze_panes = "A5"


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    ap.add_argument("--force", action="store_true", help="перезаписать существующий файл")
    args = ap.parse_args()
    if args.out.exists() and not args.force:
        sys.exit(f"{args.out} уже существует — в нём может быть ваш выбор. Используйте --force.")

    text = SOURCE.read_text(encoding="utf-8")
    rows, base = parse_metrics(text), parse_base_set(text)
    numbers = [r[0] for r in rows]
    assert len(numbers) == len(set(numbers)), "номера метрик повторяются"
    assert set(base) <= set(numbers), f"в базовом наборе есть номера не из таблиц: {set(base) - set(numbers)}"
    for number, _, options, *_ in QUESTIONS:
        assert not any("," in o for o in options) and len(",".join(options)) < 250, f"{number}: варианты ответа"

    wb = Workbook()
    build_metrics(wb.active, rows, base)
    build_questions(wb.create_sheet("Вопросы"))
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(args.out)
    status = Counter(r[6].split(":")[0] for r in rows)
    print(f"Сохранено: {args.out} — метрик {len(rows)}, в базовом наборе {len(base)}; в модели: "
          + ", ".join(f"{k} — {v}" for k, v in status.most_common()))


if __name__ == "__main__":
    main()
