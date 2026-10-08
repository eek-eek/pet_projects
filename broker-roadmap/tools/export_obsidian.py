#!/usr/bin/env python3
"""Экспорт проекта «Дорожная карта брокера» в хранилище Obsidian.

Собирает obsidian/Дорожная карта брокера/: главная заметка, этап 1, решения,
заметка на каждую метрику (формула, связи, значения модели, чувствительность),
дерево метрик (заметка и холст .canvas), отчёт с бенчмарками и модель Excel.

Значения берутся из пересчитанного broker_roadmap_model.xlsx, описания — из
tools/build_model.py. Раздел «Мои заметки» в заметках метрик сохраняется
при повторном экспорте.

Запуск: python3 tools/export_obsidian.py
"""

import json
import re
import shutil
from pathlib import Path

from openpyxl import load_workbook

import build_model as bm

ROOT = Path(__file__).resolve().parent.parent          # broker-roadmap/
REPO = ROOT.parent
VAULT = REPO / "obsidian" / "Дорожная карта брокера"
MODEL_XLSX = ROOT / "broker_roadmap_model.xlsx"
REPORT_TITLE = "Бенчмарки метрик публичных брокеров"
MY_NOTES = "## Мои заметки"

FOLDERS = {
    "result": "Метрики/1 Результат",
    "component": "Метрики/2 Источники дохода и затраты",
    "driver": "Метрики/3 Драйверы",
    "lead": "Метрики/4 Опережающие индикаторы",
    "guard": "Метрики/5 Ограничители",
    "external": "Метрики/6 Внешние факторы",
    "kpi": "Метрики/7 KPI",
}
RESULT_CODES = ["CM", "NR", "CTS", "ACQ", "LOSS"]
COMPONENT_CODES = [f"R{k}" for k in range(1, 8)] + [f"C{k}" for k in range(1, 7)]
KPI_CODES = ["ARPU", "YIELD", "CM_MARGIN", "CM_PC", "CONTRIB_PC", "LTV", "LTV_CAC", "PAYBACK", "RATE_SHARE"]
BASE_CODES = ["N_AVG", "NEW", "CHURNED", "ACTIVE", "AUC", "TURN", "MPORT", "CASHBAL", "FXVOLUME", "PRDVOLUME",
              "AUM", "SUBS"]
SEGS = ["МР", "СК", "ИК"]


# --- форматирование ---------------------------------------------------------------
def _num(v, decimals):
    s = f"{v:,.{decimals}f}".replace(",", " ").replace(".", ",")
    return s


def fmt(v, unit):
    if v is None or v == "":
        return "—"
    if unit == "%":
        return f"{v * 100:.3g}".replace(".", ",") + "%"
    if unit == "сдвиг":
        return f"{v * 100:+.3g}".replace(".", ",") + " п.п."
    if unit == "у.е.":
        if abs(v) >= 1e6:
            return _num(v / 1e6, 1) + " млн"
        return _num(v, 0 if abs(v) >= 100 or float(v).is_integer() else 2)
    if unit in ("чел.", "шт."):
        return _num(v, 0 if abs(v) >= 100 or float(v).is_integer() else 1)
    if unit in ("шт./клиент", "б.п.", "мес."):
        return _num(v, 1)
    if unit == "x":
        return _num(v, 1) + "x"
    if unit == "коэф.":
        return _num(v, 2)
    return str(v)


def money(v):
    if v is None:
        return "—"
    sign = "+" if v > 0 else ("−" if v < 0 else "")
    a = abs(v)
    if a >= 1e6:
        return f"{sign}{_num(a / 1e6, 2)} млн"
    if a >= 1e3:
        return f"{sign}{_num(a / 1e3, 0)} тыс."
    return f"{sign}{_num(a, 0)}"


def yaml_value(v):
    return json.dumps(v, ensure_ascii=False)


# --- данные модели -----------------------------------------------------------------
def load_values():
    wb = load_workbook(MODEL_XLSX, data_only=True)
    ws = wb["Модель"]
    model = {}
    for r in range(bm.MODEL_FIRST_ROW, ws.max_row + 1):
        code = ws.cell(row=r, column=1).value
        if code:
            model[code] = [ws.cell(row=r, column=c).value for c in (5, 6, 7, 8)]
    ws = wb["Чувствительность"]
    sens = {}
    for r in range(5, ws.max_row + 1):
        code = ws.cell(row=r, column=1).value
        if code:
            sens[code] = {
                "y1": [ws.cell(row=r, column=c).value for c in (10, 11, 12, 13)],
                "lt": [ws.cell(row=r, column=c).value for c in (14, 15, 16, 17)],
                "rank": ws.cell(row=r, column=19).value,
                "step": ws.cell(row=r, column=20).value,
                "step_y1": ws.cell(row=r, column=21).value,
                "step_lt": ws.cell(row=r, column=22).value,
            }
    ws = wb["Вводные"]
    globals_ = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=4).value for r in range(5, 10)}
    return model, sens, globals_


# --- связи -------------------------------------------------------------------------
ALL_CODES = (set(bm.INPUTS) | set(bm.CALC) | {e[0] for e in bm.EXTRA} | {g[0] for g in bm.GUARDRAILS}
             | {x[0] for x in bm.EXTERNAL})
CODE_RE = re.compile(r"\b[A-Z][A-Z0-9_]*[A-Z0-9]\b|\bR[1-7]\b|\bC[1-6]\b")


def codes_in(text):
    found = []
    for m in CODE_RE.findall(text or ""):
        if m in ALL_CODES and m not in found:
            found.append(m)
    return found


def build_links():
    children, parents = {}, {}
    for code, row in bm.CALC.items():
        kids = [k for k in re.findall(r"\{(\w+)\}", row[5]) if k in ALL_CODES]
        kids = list(dict.fromkeys(kids))
        children[code] = kids
        for k in kids:
            parents.setdefault(k, []).append(code)
    reverse = {"lead": {}, "guard": {}, "external": {}}
    for e in bm.EXTRA:
        for c in codes_in(e[8]):
            reverse["lead"].setdefault(c, []).append(e[0])
    for g in bm.GUARDRAILS:
        for c in codes_in(g[6]):
            reverse["guard"].setdefault(c, []).append(g[0])
    for x in bm.EXTERNAL:
        for c in codes_in(x[4]):
            reverse["external"].setdefault(c, []).append(x[0])
    return children, parents, reverse


def wl(code):
    return f"[[{code}]]"


def name_of(code):
    if code in bm.INPUTS or code in bm.CALC:
        return bm.meta(code)[0]
    for e in bm.EXTRA:
        if e[0] == code:
            return e[1]
    for g in bm.GUARDRAILS:
        if g[0] == code:
            return g[1]
    for x in bm.EXTERNAL:
        if x[0] == code:
            return x[1]
    return code


def link_list(codes):
    return ", ".join(f"{wl(c)} ({name_of(c)})" for c in codes) if codes else "—"


# --- запись заметок ----------------------------------------------------------------
def keep_my_notes(path):
    if path.exists():
        text = path.read_text(encoding="utf-8")
        if MY_NOTES in text:
            return text[text.index(MY_NOTES):].rstrip() + "\n"
    return MY_NOTES + "\n\n"


def write_note(folder, code, front, body):
    path = VAULT / folder / f"{code}.md"
    path.parent.mkdir(parents=True, exist_ok=True)
    tail = keep_my_notes(path)
    fm = "---\n" + "\n".join(f"{k}: {yaml_value(v)}" for k, v in front.items()) + "\n---\n"
    path.write_text(fm + body.rstrip() + "\n\n" + tail, encoding="utf-8")


def values_table(code, unit, model):
    vals = model.get(code)
    if not vals:
        return ""
    head = "| МР | СК | ИК | Итого |\n|---:|---:|---:|---:|\n"
    total = vals[3] if vals[3] is not None else None
    return head + "| " + " | ".join(fmt(v, unit) for v in vals[:3]) + f" | {fmt(total, unit)} |\n"


def sens_block(code, sens):
    s = sens.get(code)
    if not s:
        return ""
    lines = ["## Чувствительность маржинального дохода", "",
             "Сколько у.е. МД в год даёт улучшение драйвера при прочих равных (данные условные).", "",
             "| Улучшение | МР | СК | ИК | Итого |", "|---|---:|---:|---:|---:|",
             "| +1%, год 1 | " + " | ".join(money(v) for v in s["y1"]) + " |",
             "| +1%, долгосрочно | " + " | ".join(money(v) for v in s["lt"]) + " |"]
    if s["step"]:
        lines.append(f"| {s['step']}, год 1 | | | | {money(s['step_y1'])} |")
        lines.append(f"| {s['step']}, долгосрочно | | | | {money(s['step_lt'])} |")
    if s["rank"] not in (None, ""):
        lines += ["", f"Ранг в своей группе по долгосрочному эффекту: {s['rank']}."]
    return "\n".join(lines) + "\n"


def export_model_metrics(model, sens, children, parents, reverse):
    for code in RESULT_CODES + COMPONENT_CODES + BASE_CODES + KPI_CODES + list(bm.INPUTS):
        name, unit, ftext = bm.meta(code)
        if code in bm.INPUTS:
            i = bm.INPUTS[code]
            src, freq, owner, ll, risk = bm.P_INPUT[code]
            folder, level, typ = FOLDERS["driver"], "L3", "Драйвер (вводная)"
            definition = bm.EXPLAIN[code] + (f" {i.comment}" if i.comment else "")
            extra = {"Направление": "больше — лучше" if i.direction > 0 else "меньше — лучше",
                     "Управляемость": i.control, "Тип рычага": i.lever}
            benchmark = i.source
        else:
            lvl, typ, grp, definition, src, freq, owner, ll, risk = bm.P[code]
            definition = bm.EXPLAIN.get(code, definition)
            level = lvl
            folder = (FOLDERS["result"] if code in RESULT_CODES else FOLDERS["component"] if code in COMPONENT_CODES
                      else FOLDERS["kpi"] if code in KPI_CODES else FOLDERS["driver"])
            extra, benchmark = {}, ""
        tree_parents = [p for p in parents.get(code, []) if p not in KPI_CODES]
        kpi_users = [p for p in parents.get(code, []) if p in KPI_CODES]
        front = {"код": code, "название": name, "уровень": level, "тип": typ, "единица": unit,
                 "частота": freq, "владелец": owner, "tags": ["метрика", level, typ.split()[0].lower()],
                 "aliases": [name]}
        body = [f"# {code} — {name}", "", f"**Что это.** {definition}", "",
                f"**Формула.** {ftext}" + ("" if code in bm.INPUTS else f"  \n`{bm.CALC[code][5]}`"), ""]
        if code in children and children[code]:
            body += [f"**Состоит из:** {link_list(children[code])}", ""]
        if tree_parents:
            body += [f"**Входит в:** {link_list(tree_parents)}", ""]
        if kpi_users:
            body += [f"**Используется в KPI:** {link_list(kpi_users)}", ""]
        for key, title in [("lead", "Опережающие индикаторы"), ("guard", "Ограничители"),
                           ("external", "Внешние факторы")]:
            if reverse[key].get(code):
                body += [f"**{title}:** {link_list(reverse[key][code])}", ""]
        if extra:
            body += [" · ".join(f"**{k}:** {v}" for k, v in extra.items()), ""]
        body += ["## Значения в модели (условные)", "", values_table(code, unit, model)]
        if code in bm.INPUTS and benchmark:
            body += ["**Ориентир:** " + benchmark, ""]
        sb = sens_block(code, sens)
        if sb:
            body += [sb]
        body += ["## Паспорт", "", f"- Опережающая или запаздывающая: {ll}", f"- Источник данных: {src}",
                 f"- Частота: {freq}", f"- Владелец: {owner}"]
        if risk:
            body += [f"- Риски интерпретации: {risk}"]
        write_note(folder, code, front, "\n".join(body))


def export_other_metrics():
    for e in bm.EXTRA:
        code, name, grp, definition, formula, unit, segs, d, link, src, freq, owner = e
        front = {"код": code, "название": name, "уровень": "L4", "тип": "Опережающий индикатор", "группа": grp,
                 "единица": unit, "частота": freq, "владелец": owner,
                 "tags": ["метрика", "L4", "индикатор"], "aliases": [name]}
        body = [f"# {code} — {name}", "", f"**Что это.** {definition}.", "", f"**Как считать.** {formula}", "",
                f"**Двигает драйверы:** {link_list(codes_in(link))}", "",
                f"**Группа:** {grp} · **Сегменты:** {segs} · **Направление:** {d}", "",
                "## Паспорт", "", f"- Источник данных: {src}", f"- Частота: {freq}", f"- Владелец: {owner}"]
        write_note(FOLDERS["lead"], code, front, "\n".join(body))
    for g in bm.GUARDRAILS:
        code, name, definition, formula, unit, d, link, src, freq, owner = g
        front = {"код": code, "название": name, "уровень": "G", "тип": "Ограничитель", "единица": unit,
                 "частота": freq, "владелец": owner, "tags": ["метрика", "ограничитель"], "aliases": [name]}
        body = [f"# {code} — {name}", "", f"**Что это.** {definition}.", "", f"**Как считать.** {formula}", "",
                f"**Связан с:** {link_list(codes_in(link))}", "", f"**Направление:** {d}", "",
                "Гипотеза принимается, только если ограничитель остаётся в норме.", "",
                "## Паспорт", "", f"- Источник данных: {src}", f"- Частота: {freq}", f"- Владелец: {owner}"]
        write_note(FOLDERS["guard"], code, front, "\n".join(body))
    for x in bm.EXTERNAL:
        code, name, definition, unit, link, src, freq, owner = x
        front = {"код": code, "название": name, "уровень": "X", "тип": "Внешний фактор", "единица": unit,
                 "частота": freq, "владелец": owner, "tags": ["метрика", "внешний-фактор"], "aliases": [name]}
        body = [f"# {code} — {name}", "", f"**Что это.** {definition}.", "",
                f"**Влияет на:** {link_list(codes_in(link)) if codes_in(link) else link}", "",
                "Гипотезы на внешние факторы не претендуют: они задают сценарии и нормируют результат на рынок.", "",
                "## Паспорт", "", f"- Источник данных: {src}", f"- Частота: {freq}", f"- Владелец: {owner}"]
        write_note(FOLDERS["external"], code, front, "\n".join(body))


def export_tree_note():
    lines = ["# Дерево метрик", "",
             "Финансовое дерево от маржинального дохода к драйверам. Визуально — на холсте [[Дерево метрик.canvas]].",
             ""]
    for item in bm.TREE:
        if item[0] == "section":
            lines += ["", f"## {item[1]}", ""]
            continue
        level, code = item
        lines.append("  " * level + f"- {wl(code)} — {bm.meta(code)[0]}")
    for title, rows in [("Опережающие индикаторы (L4)", [(e[0], e[1], e[2]) for e in bm.EXTRA]),
                        ("Ограничители", [(g[0], g[1], "") for g in bm.GUARDRAILS]),
                        ("Внешние факторы", [(x[0], x[1], "") for x in bm.EXTERNAL])]:
        lines += ["", f"## {title}", ""]
        group = None
        for code, name, grp in rows:
            if grp and grp != group:
                lines += ["", f"**{grp}**", ""]
                group = grp
            lines.append(f"- {wl(code)} — {name}")
    path = VAULT / "Дерево метрик.md"
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def export_canvas():
    nodes, edges, stack = [], [], {}
    y = 0
    for idx, item in enumerate(bm.TREE):
        if item[0] == "section":
            y += 60
            nodes.append({"id": f"s{idx}", "type": "text", "text": f"## {item[1]}", "x": 0, "y": y,
                          "width": 900, "height": 70})
            y += 110
            stack = {}
            continue
        level, code = item
        node_id = f"n{idx}"
        color = ("4" if code in RESULT_CODES else "5" if code.startswith("R") and code in COMPONENT_CODES
                 else "2" if code in COMPONENT_CODES else "3" if code in bm.INPUTS
                 and bm.INPUTS[code].control == "Внешний" else "6" if code in bm.INPUTS
                 and bm.INPUTS[code].lever == "Ценовой" else None)
        node = {"id": node_id, "type": "text", "text": f"**[[{code}]]**\n{bm.meta(code)[0]}",
                "x": level * 440, "y": y, "width": 380, "height": 90}
        if color:
            node["color"] = color
        nodes.append(node)
        if level - 1 in stack:
            edges.append({"id": f"e{idx}", "fromNode": stack[level - 1], "fromSide": "right",
                          "toNode": node_id, "toSide": "left"})
        stack[level] = node_id
        for deeper in [k for k in stack if k > level]:
            del stack[deeper]
        y += 110
    legend = ("**Цвета**\n\nзелёный — результат\nголубой — источники дохода\nоранжевый — затраты и потери\n"
              "фиолетовый — ценовые рычаги\nжёлтый — внешние факторы")
    nodes.append({"id": "legend", "type": "text", "text": legend, "x": -460, "y": 0, "width": 380, "height": 220})
    (VAULT / "Дерево метрик.canvas").write_text(
        json.dumps({"nodes": nodes, "edges": edges}, ensure_ascii=False, indent=1), encoding="utf-8")


def link_codes_in_markdown(text):
    """`КОД` → [[КОД]] вне блоков кода для кодов, у которых есть заметка."""
    out, in_fence = [], False
    for line in text.splitlines():
        if line.strip().startswith("```"):
            in_fence = not in_fence
            out.append(line)
            continue
        if not in_fence:
            line = re.sub(r"`([A-Z][A-Z0-9_]*)`", lambda m: f"[[{m.group(1)}]]" if m.group(1) in ALL_CODES
                          else m.group(0), line)
        out.append(line)
    return "\n".join(out) + "\n"


def to_vault_markdown(text):
    """Ссылки репозитория → ссылки Obsidian, `КОД` → [[КОД]]."""
    text = text.replace("[broker_roadmap_model.xlsx](broker_roadmap_model.xlsx)", "[[broker_roadmap_model.xlsx]]")
    text = re.sub(r"\[([^\]]+)\]\(\.\./reports/[^)]+\)", lambda m: (
        f"[[{REPORT_TITLE}]]" if m.group(1).startswith("reports/") else f"[[{REPORT_TITLE}|{m.group(1)}]]"), text)
    return link_codes_in_markdown(text)


def export_documents():
    for source, title in [("01_metrics.md", "Этап 1 — метрики"),
                          ("01b_best_practice_metrics.md", "Метрики — лучшие практики")]:
        text = (ROOT / source).read_text(encoding="utf-8")
        (VAULT / f"{title}.md").write_text(to_vault_markdown(text), encoding="utf-8")

    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    decisions = [line for line in readme.splitlines() if re.match(r"^\| \d{2}\.\d{2}\.\d{4} \|", line)]
    (VAULT / "Решения.md").write_text(
        "# Решения\n\nЖурнал решений по проекту.\n\n| Дата | Решение |\n|---|---|\n" + "\n".join(decisions) + "\n",
        encoding="utf-8")

    bench = VAULT / "Бенчмарки"
    (bench / "Заметки исследования").mkdir(parents=True, exist_ok=True)
    shutil.copy(REPO / "reports" / f"{REPORT_TITLE}.md", bench / f"{REPORT_TITLE}.md")
    for f in (REPO / "research_notes" / REPORT_TITLE).glob("*.md"):
        shutil.copy(f, bench / "Заметки исследования" / f.name)
    (VAULT / "Модель").mkdir(exist_ok=True)
    shutil.copy(MODEL_XLSX, VAULT / "Модель" / MODEL_XLSX.name)


# --- справочник метрик одной страницей ---------------------------------------------
SHORT = {
    "CM": "Маржинальный доход", "NR": "Чистая выручка", "CTS": "Затраты на обслуживание",
    "ACQ": "Затраты на привлечение", "LOSS": "Потери", "R1": "Торговые комиссии", "R2": "Маржинальное кредитование",
    "R3": "Доход на свободные остатки", "R4": "Конвертация валют", "R5": "Размещения и продукты",
    "R6": "ДУ и консультирование", "R7": "Подписки и сервисные комиссии", "C1": "Поддержка клиентов",
    "C2": "Онбординг и KYC", "C3": "Ручные операции бэк-офиса", "C4": "Менеджеры и покрытие",
    "C5": "Кредитные потери", "C6": "Операционные потери", "N0": "Клиенты на начало", "LEADS": "Заявки",
    "CR_OPEN": "Конверсия в счёт", "CR_FUND": "Конверсия в фондирование", "NEW": "Новые клиенты", "CHURN": "Отток",
    "CHURNED": "Ушедшие клиенты", "N_AVG": "Фондированные клиенты", "ACT": "Доля активных",
    "ACTIVE": "Активные клиенты", "TPA": "Сделок на активного в месяц", "TICKET": "Средний чек сделки",
    "TURN": "Оборот", "TAKE": "Комиссия брокера", "EXEC": "Затраты на исполнение", "AUC_PC": "Активы на клиента",
    "AUC": "Активы клиентов", "CASH": "Доля свободных денег", "CASHBAL": "Свободные остатки",
    "FLOAT_Y": "Доходность размещения", "FLOAT_P": "Ставка клиентам на остатки", "MRG_PEN": "Проникновение маржи",
    "MRG_DEBT": "Задолженность на клиента", "MPORT": "Маржинальный портфель", "MRG_RATE": "Ставка по марже",
    "FUND_RATE": "Стоимость фондирования", "CL": "Уровень кредитных потерь", "FX_PEN": "Доля конвертирующих",
    "FX_VOL": "Конвертации на клиента", "FXVOLUME": "Объём конвертаций", "FX_SPREAD": "Спред конвертации",
    "PRD_PEN": "Проникновение продуктов", "PRD_VOL": "Покупки на клиента", "PRDVOLUME": "Объём продаж продуктов",
    "PRD_FEE": "Комиссия по продуктам", "AM_SHARE": "Доля активов в ДУ", "AUM": "Активы в ДУ",
    "AM_FEE": "Комиссия за управление", "SUB_PEN": "Проникновение подписки", "SUBS": "Подписчики",
    "SUB_PRICE": "Цена подписки в год", "CUST_FEE": "Кастодиальная комиссия", "SUP_RATE": "Обращений на клиента в год",
    "SUP_COST": "Стоимость обращения", "KYC_UNIT": "Авто-проверка заявки", "KYC_MAN": "Доля ручной проверки",
    "KYC_MAN_COST": "Ручная проверка заявки", "OPS_PC": "Операций на клиента в год",
    "OPS_MAN": "Доля ручных операций", "OPS_COST": "Стоимость ручной операции", "RM_LOAD": "Клиентов на менеджера",
    "RM_COST": "Стоимость менеджера в год", "CAC": "Стоимость привлечения клиента", "OPLOSS": "Операционные потери",
    "ARPU": "Выручка на клиента, у.е.", "YIELD": "Доходность активов, б.п.", "CM_MARGIN": "Маржинальность",
    "CM_PC": "МД на клиента, у.е.", "CONTRIB_PC": "Вклад клиента в год, у.е.", "LTV": "Ценность клиента (LTV), у.е.",
    "LTV_CAC": "LTV ÷ CAC", "PAYBACK": "Окупаемость привлечения, мес.", "RATE_SHARE": "Доля процентного дохода",
}
FORM = {
    "CM": "ЧВ − ЗО − ЗП − ПТ", "NR": "R1 + … + R7", "CTS": "C1 + C2 + C3 + C4", "ACQ": "новые клиенты × CAC",
    "LOSS": "C5 + C6", "R1": "оборот × (комиссия − затраты на исполнение)",
    "R2": "маржинальный портфель × (ставка по марже − фондирование)",
    "R3": "свободные остатки × (доходность размещения − ставка клиентам)", "R4": "объём конвертаций × спред",
    "R5": "объём продаж × комиссия", "R6": "активы в ДУ × комиссия за управление",
    "R7": "подписчики × цена + активы × кастодиальная комиссия", "C1": "клиенты × обращения × стоимость обращения",
    "C2": "заявки × (авто-проверка + доля ручной × стоимость ручной)",
    "C3": "клиенты × операции × доля ручных × стоимость ручной", "C4": "клиенты ÷ нагрузка × стоимость менеджера",
    "C5": "маржинальный портфель × уровень потерь", "C6": "ошибки, компенсации, штрафы",
    "NEW": "заявки × конверсия в счёт × конверсия в фондирование", "CHURNED": "клиенты на начало × отток",
    "N_AVG": "клиенты на начало + ½ (новые − ушедшие)", "ACTIVE": "фондированные × доля активных",
    "TURN": "активные × сделки × 12 × средний чек", "AUC": "фондированные × активы на клиента",
    "CASHBAL": "активы × доля свободных денег", "MPORT": "фондированные × проникновение маржи × задолженность",
    "FXVOLUME": "фондированные × доля конвертирующих × конвертации на клиента",
    "PRDVOLUME": "фондированные × проникновение продуктов × покупки на клиента", "AUM": "активы × доля в ДУ",
    "SUBS": "фондированные × проникновение подписки",
}
DRIVER_GROUPS = [
    ("Клиентская база и воронка", ["N0", "LEADS", "CR_OPEN", "CR_FUND", "NEW", "CHURN", "CHURNED", "N_AVG", "ACT",
                                   "ACTIVE"]),
    ("Торговля → R1", ["TPA", "TICKET", "TURN", "TAKE", "EXEC"]),
    ("Активы и свободные остатки → R3, R6, R7", ["AUC_PC", "AUC", "CASH", "CASHBAL", "FLOAT_Y", "FLOAT_P"]),
    ("Маржинальное кредитование → R2, C5", ["MRG_PEN", "MRG_DEBT", "MPORT", "MRG_RATE", "FUND_RATE", "CL"]),
    ("Валюта → R4", ["FX_PEN", "FX_VOL", "FXVOLUME", "FX_SPREAD"]),
    ("Размещения и продукты → R5", ["PRD_PEN", "PRD_VOL", "PRDVOLUME", "PRD_FEE"]),
    ("ДУ и консультирование → R6", ["AM_SHARE", "AUM", "AM_FEE"]),
    ("Подписки и сервисы → R7", ["SUB_PEN", "SUBS", "SUB_PRICE", "CUST_FEE"]),
    ("Затраты на обслуживание → C1–C4", ["SUP_RATE", "SUP_COST", "KYC_UNIT", "KYC_MAN", "KYC_MAN_COST", "OPS_PC",
                                         "OPS_MAN", "OPS_COST", "RM_LOAD", "RM_COST"]),
    ("Привлечение и потери → ЗП, C6", ["CAC", "OPLOSS"]),
]


def _trim(x, d=1):
    s = f"{x:,.{d}f}".replace(",", " ").replace(".", ",")
    return s.rstrip("0").rstrip(",") if "," in s else s


def short_fmt(v, unit, code=""):
    if v is None or (v == 0 and code in ("RM_LOAD", "RM_COST", "SUB_PRICE")):
        return "—"
    if unit == "%":
        return f"{v * 100:.3g}".replace(".", ",") + "%"
    if unit == "x":
        return _trim(v) + "x"
    a = abs(v)
    if a >= 1e9:
        return _trim(v / 1e9) + " млрд"
    if a >= 1e6:
        return _trim(v / 1e6) + " млн"
    if a >= 10 and unit in ("у.е.", "чел.", "шт."):
        return _trim(v, 0)
    return _trim(v)


def reference_markdown(model, links):
    """Справочник метрик: links=True — коды как [[ссылки]] (Obsidian), иначе `код` (чат)."""
    ref = (lambda c: f"[[{c}]]") if links else (lambda c: f"`{c}`")
    seg = lambda c: " / ".join(short_fmt(x, bm.meta(c)[1], c) for x in model[c][:3])
    mln = lambda c: " / ".join(_trim(v / 1e6) for v in model[c][:3])
    out = ["### Уровень 0–1. Результат", "",
           "| Код | Метрика = формула | Что это | МР / СК / ИК, млн у.е. |", "|---|---|---|---|"]
    out += [f"| {ref(c)} | **{SHORT[c]}** = {FORM[c]} | {bm.EXPLAIN[c]} | {mln(c)} |" for c in RESULT_CODES]
    for title, codes in [("Уровень 2. Источники дохода", COMPONENT_CODES[:7]),
                         ("Уровень 2. Затраты и потери", COMPONENT_CODES[7:])]:
        out += ["", f"### {title}", "", "| Код | Метрика = формула | Что это | МР / СК / ИК, млн у.е. |",
                "|---|---|---|---|"]
        out += [f"| {ref(c)} | **{SHORT[c]}** = {FORM[c]} | {bm.EXPLAIN[c]} | {mln(c)} |" for c in codes]
    out += ["", "### Уровень 3. Драйверы", "",
            "Значения условные, деньги — в у.е. *(расчёт)* — считается из других драйверов, *(цена)* — ценовой "
            "рычаг, *(внешний)* — задаётся рынком."]
    for title, codes in DRIVER_GROUPS:
        out += ["", f"**{title}**", "", "| Код | Драйвер | Что это | МР / СК / ИК |", "|---|---|---|---|"]
        for c in codes:
            tag, expl = "", bm.EXPLAIN[c]
            if c in bm.CALC:
                tag, expl = " *(расчёт)*", f"= {FORM[c]}. {expl}"
            elif bm.INPUTS[c].lever == "Ценовой":
                tag = " *(цена)*"
            elif bm.INPUTS[c].control == "Внешний":
                tag = " *(внешний)*"
            out.append(f"| {ref(c)} | {SHORT[c]}{tag} | {expl} | {seg(c)} |")
    out += ["", "### Уровень 4. Опережающие индикаторы", "",
            "То, что гипотезы двигают напрямую и что видно через недели; стрелка — на какой драйвер влияет.", ""]
    order, by = [], {}
    for e in bm.EXTRA:
        if e[2] not in by:
            order.append(e[2])
            by[e[2]] = []
        by[e[2]].append(f"{ref(e[0])} {e[1]} → {e[8]}")
    out += [f"- **{g}:** " + "; ".join(by[g]) + "." for g in order]
    out += ["", "### Ограничители", "", "Не должны ухудшаться при реализации гипотез.", ""]
    out += [f"- {ref(g[0])} **{g[1]}** — {g[2][0].lower() + g[2][1:]}." for g in bm.GUARDRAILS]
    out += ["", "### Внешние факторы", "", "Задают сценарии; гипотезы на них не претендуют.", ""]
    ext_link = {"X_REG": "многие драйверы", "X_SHARE": "нормирует оценку всех гипотез на рынок"}
    out += [f"- {ref(x[0])} **{x[1]}** → {ext_link.get(x[0], x[4])}." for x in bm.EXTERNAL]
    out += ["", "### Производные KPI", "", "| Код | KPI | Что это | МР / СК / ИК |", "|---|---|---|---|"]
    out += [f"| {ref(c)} | {SHORT[c]} | {bm.EXPLAIN[c]} | {seg(c)} |" for c in KPI_CODES]
    return "\n".join(out) + "\n"


def export_reference(model):
    text = ("# Справочник метрик\n\nВсе метрики дерева на одной странице: что это, как считается, условные "
            "значения по сегментам. Подробности — в заметке каждой метрики, структура — [[Дерево метрик]].\n\n")
    (VAULT / "Справочник метрик.md").write_text(text + reference_markdown(model, links=True), encoding="utf-8")


def export_home(model, globals_):
    cm, nr = model["CM"], model["NR"]
    text = f"""# Дорожная карта брокера

> [!info] Статус
> Этап 1 «Метрики» — в работе: решения внесены, метрики на согласовании.

Цель — дорожная карта доработок и продуктовых инициатив брокера, где каждый проект связан с измеримым эффектом на маржинальный доход.

## Этапы

| Этап | Результат | Статус |
|---|---|---|
| 1. Метрики | Дерево метрик, паспорта, модель чувствительности — [[Этап 1 — метрики]] | В работе |
| 2. Гипотезы | Шаблон карты гипотез, привязка к драйверам | — |
| 3. Трудозатраты | Оценка гипотез по ролям и ресурсам | — |
| 4. Дорожная карта | Приоритизация и раскладка на 3 года по полугодиям | — |

## С чего начать

1. [[Метрики — лучшие практики]] — стандартный набор метрик отрасли: что раскрывают ведущие брокеры.
2. [[Этап 1 — метрики]] — главный документ этапа: решения и выводы.
3. [[Справочник метрик]] — метрики модели на одной странице с пояснениями и значениями.
4. [[Дерево метрик]] — структура модели по уровням; визуально — [[Дерево метрик.canvas]].
5. [[Решения]] — журнал решений.
6. [[{REPORT_TITLE}]] — ориентиры по рынку, справочно.
7. Модель с расчётами — [[broker_roadmap_model.xlsx]], начните с листа «Сводка».

## Ключевые цифры модели (условные, год 1)

| | МР | СК | ИК | Итого |
|---|---:|---:|---:|---:|
| [[NR]] Чистая выручка, у.е. | {' | '.join(fmt(v, 'у.е.') for v in nr)} |
| [[CM]] Маржинальный доход, у.е. | {' | '.join(fmt(v, 'у.е.') for v in cm)} |

Сценарий ставок: {fmt(globals_.get('RATE_SHIFT'), 'сдвиг')} к текущим.

## Как устроено хранилище

- Заметки метрик (папка «Метрики») собраны из модели скриптом `broker-roadmap/tools/export_obsidian.py` в репозитории. При повторном экспорте они пересоздаются, но раздел «Мои заметки» внизу каждой заметки сохраняется — пишите комментарии туда.
- Цифры живут в модели Excel; документы этапов — в заметках этапов.
"""
    (VAULT / "Дорожная карта брокера.md").write_text(text, encoding="utf-8")


def main():
    VAULT.mkdir(parents=True, exist_ok=True)
    model, sens, globals_ = load_values()
    children, parents, reverse = build_links()
    export_model_metrics(model, sens, children, parents, reverse)
    export_other_metrics()
    export_tree_note()
    export_reference(model)
    export_canvas()
    export_documents()
    export_home(model, globals_)
    notes = list(VAULT.rglob("*.md"))
    print(f"Хранилище: {VAULT} — заметок {len(notes)}")


if __name__ == "__main__":
    main()
